#!/usr/bin/env python3
"""Convert the slides PDF into a PPTX with one full-page image per slide.

Usage: pdf2pptx.py input.pdf output.pptx [dpi]

If input.notes exists (written by the slides build, see \\nsl in
mathmacros-sl.tex), its script-only paragraphs become the speaker notes of
their slides, as plain text.

Needs pdftoppm (poppler-utils), python-pptx and pylatexenc; the Makefile runs
it from the venv in ~/.local/share/lecture-tools/venv.
"""
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu
from pylatexenc.latex2text import LatexNodes2Text, MacroTextSpec, get_default_latex_context_db

src, dst = Path(sys.argv[1]), Path(sys.argv[2])
dpi = sys.argv[3] if len(sys.argv) > 3 else "200"

info = subprocess.run(["pdfinfo", src], capture_output=True, text=True, check=True).stdout
w_pt, h_pt = map(float, re.search(r"Page size:\s+([\d.]+) x ([\d.]+) pts", info).groups())

# notes file: one line "page|detokenized LaTeX" per \\nsl paragraph
notes = {}
notes_file = src.with_suffix(".notes")
if notes_file.exists():
    db = get_default_latex_context_db()
    db.add_context_category("lecture", prepend=True,
                            macros=[MacroTextSpec(m, "%(1)s") for m in ("texttt", "url", "textsf", "textrm")])
    to_text = LatexNodes2Text(latex_context=db, math_mode="text")
    for line in notes_file.read_text(encoding="utf-8").splitlines():
        page, _, tex = line.partition("|")
        text = re.sub(r"[ \t]+", " ", to_text.latex_to_text(tex)).strip()
        notes.setdefault(int(page), []).append(text)

prs = Presentation()
prs.slide_width = Emu(round(w_pt * 12700))
prs.slide_height = Emu(round(h_pt * 12700))
blank = prs.slide_layouts[6]

with tempfile.TemporaryDirectory() as tmp:
    subprocess.run(["pdftoppm", "-r", dpi, "-png", src, f"{tmp}/p"], check=True)
    for png in sorted(Path(tmp).glob("p-*.png"), key=lambda p: int(p.stem.split("-")[1])):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
        page = int(png.stem.split("-")[1])
        if page in notes:
            slide.notes_slide.notes_text_frame.text = "\n\n".join(notes[page])

prs.save(dst)
print(f"{dst}: {len(prs.slides)} slides, {len(notes)} with notes")
